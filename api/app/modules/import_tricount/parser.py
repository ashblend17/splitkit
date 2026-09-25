"""Reads a Tricount export (CSV or XLSX) into rows. Pure: bytes in, dataclasses out.

Tricount's own "Export" writes one row per transaction with a payer column and one column
per person holding that person's share. Third-party exporters and older app versions name
the columns differently, so headers are matched case-insensitively against aliases:

    Title, Amount, Currency, Exchange rate, Amount in default currency, Date & time,
    Paid by, Transaction type, Category, Impacted to <person>…

Per-person columns may be "Impacted to X", "Paid for X", "Share X" or "For X" (the share),
and "Paid by X" (what X advanced; the Sesterce layout, used when there is no payer column).
Transaction types: Normal/Expense, Money transfer/Balance/Reimbursement, Income.

Amounts stay Decimal here; `plan.py` turns them into minor units.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

MAX_BYTES = 5 * 1024 * 1024


class ParseError(Exception):
    """The file as a whole can't be read. The message is human copy."""


# ---- header aliases ----

_COLUMNS: dict[str, set[str]] = {
    "title": {"title", "description", "what", "expense", "titre", "label", "subject"},
    "amount": {"amount", "total", "amount paid", "price", "montant", "sum", "value"},
    "currency": {"currency", "devise", "currency code"},
    "rate": {"exchange rate", "rate", "conversion rate", "taux de change"},
    "amount_default": {
        "amount in default currency",
        "amount in main currency",
        "amount in tricount currency",
        "amount in base currency",
        "default currency amount",
        "converted amount",
    },
    "date": {
        "date",
        "date & time",
        "date and time",
        "date time",
        "datetime",
        "transaction date",
        "day",
    },
    "payer": {"paid by", "who paid", "payer", "paye par", "from"},
    "type": {"transaction type", "type", "kind"},
    "category": {"category", "categorie"},
    "custom_category": {"custom category"},
}
# Longest first, so "impacted to " wins over "impacted ".
_SHARE_PREFIXES = ("impacted to ", "impacted for ", "paid for ", "impacted ", "share ", "for ")
_PAID_BY_PREFIX = "paid by "
# Columns other exporters add next to the shares; never a person.
_IGNORED_PREFIXES = ("share ratio ", "local share ", "allocation type ", "local ")

_KINDS: dict[str, set[str]] = {
    "expense": {"normal", "expense", "expenses", "depense", "spending", "purchase"},
    "transfer": {
        "money transfer",
        "transfer",
        "balance",
        "reimbursement",
        "payment",
        "settlement",
        "remboursement",
        "transfert d'argent",
    },
    "income": {"income", "revenu", "recette", "money received"},
}
_SYMBOLS = {"₹": "INR", "$": "USD", "€": "EUR", "£": "GBP", "rs": "INR", "rs.": "INR"}


def _norm(text: Any) -> str:
    """Lowercase, no accents, no BOM, single spaces: "Catégorie " -> "categorie"."""
    s = unicodedata.normalize("NFKD", str(text or "").replace("﻿", ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s.replace("_", " ")).strip().lower().rstrip(":").strip()


def _base(header: str) -> str:
    """Drop a trailing "(EUR)" so "Amount in default currency (EUR)" matches."""
    return re.sub(r"\s*\([^)]*\)$", "", header).strip()


# ---- result types ----


@dataclass
class RawRow:
    line: int
    """1-based row number in the file, for "Row 12" in messages."""
    title: str
    date: date | None
    date_text: str
    kind: str
    """expense | transfer | income"""
    kind_text: str
    currency: str | None
    amount: Decimal | None
    amount_default: Decimal | None
    payer: str | None
    paid_by: dict[str, Decimal]
    shares: dict[str, Decimal]
    """Person -> share in the row's currency. Blank cells are left out."""
    category: str | None
    problems: list[str] = field(default_factory=list)
    """Cells that couldn't be read; each one makes the row unusable."""
    notes: list[str] = field(default_factory=list)
    """Things worth a look that don't stop the row."""


@dataclass
class ParsedFile:
    rows: list[RawRow]
    people: list[str]
    """Everyone in the file, in column order, then payers not seen in a column."""
    title: str | None
    """A tricount name found above the header, if any."""


# ---- reading the table ----


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises a zoo of exceptions for a bad file
        raise ParseError(
            "This Excel file couldn't be opened. Export it from Tricount again."
        ) from e
    sheet = book.worksheets[0]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    book.close()
    return rows


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = content.decode(encoding)
            if "\x00" not in text:
                return text
        except UnicodeDecodeError:
            continue
    return content.decode("cp1252", errors="replace")


def _read_csv(content: bytes) -> list[list[Any]]:
    text = _decode(content)
    sample = "\n".join(text.splitlines()[:20])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample else ""
        delimiter = max(",;\t|", key=first.count)
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def read_table(content: bytes, file_name: str) -> list[list[Any]]:
    if not content.strip():
        raise ParseError("This file is empty.")
    if len(content) > MAX_BYTES:
        raise ParseError("This file is larger than 5 MB. Tricount exports are much smaller.")
    if content[:4] == b"PK\x03\x04" or file_name.lower().endswith((".xlsx", ".xlsm")):
        return _read_xlsx(content)
    if file_name.lower().endswith(".xls"):
        raise ParseError("Old .xls files aren't supported. Export as CSV or XLSX instead.")
    return _read_csv(content)


# ---- columns ----


@dataclass
class _Columns:
    fixed: dict[str, int]
    shares: list[tuple[int, str]]
    paid_by: list[tuple[int, str]]


def _columns(header: list[Any]) -> _Columns:
    fixed: dict[str, int] = {}
    shares: list[tuple[int, str]] = []
    paid_by: list[tuple[int, str]] = []
    for i, cell in enumerate(header):
        raw = str(cell or "").replace("﻿", "").strip()
        h = _norm(raw)
        if not h:
            continue
        key = next((k for k, names in _COLUMNS.items() if _base(h) in names), None)
        if key is not None:
            fixed.setdefault(key, i)
            continue
        if h.startswith(_IGNORED_PREFIXES):
            continue
        if h.startswith(_PAID_BY_PREFIX):
            paid_by.append((i, raw[len(_PAID_BY_PREFIX) :].strip()))
            continue
        prefix = next((p for p in _SHARE_PREFIXES if h.startswith(p)), None)
        if prefix is not None:
            shares.append((i, raw[len(prefix) :].strip()))
    return _Columns(fixed, shares, paid_by)


def _find_header(table: list[list[Any]]) -> tuple[int, _Columns]:
    for i, row in enumerate(table[:15]):
        cols = _columns(row)
        if "title" in cols.fixed and (
            "amount" in cols.fixed or cols.shares or cols.paid_by or "payer" in cols.fixed
        ):
            return i, cols
    raise ParseError(
        "We couldn't find Tricount's columns (Title, Amount, Paid by…). "
        "In Tricount, open the tricount, then Settings → Export, and upload that file."
    )


# ---- cells ----

_NUMBER_CHARS = re.compile(r"[^0-9,.\-+()]")


def decimal_style(values: list[Any]) -> bool | None:
    """True if the file writes decimal commas ("12,50"), False for dots, None if unclear."""
    comma = dot = 0
    for v in values:
        if not isinstance(v, str):
            continue
        s = _NUMBER_CHARS.sub("", v)
        if re.search(r"\d,\d{1,2}$", s) or re.search(r"\.\d{3},\d+$", s):
            comma += 1
        elif re.search(r"\d\.\d{1,2}$", s) or re.search(r",\d{3}\.\d+$", s):
            dot += 1
    if comma == dot:
        return None
    return comma > dot


def parse_amount(value: Any, decimal_comma: bool | None = None) -> Decimal | None:
    """A cell -> Decimal, or None if blank. Handles "1.234,56", "1,234.56", "1 234,56",
    "₹1,50,000", "(12.50)" and "-12,5". Raises ValueError for anything else."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | Decimal):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(repr(value))  # shortest round-trip form: 1234.56, never 1234.5599…
    text = str(value).strip()
    if not text:
        return None
    s = _NUMBER_CHARS.sub("", text.replace(" ", "").replace("'", ""))
    negative = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if not re.search(r"\d", s):
        raise ValueError(text)
    if "," in s and "." in s:
        decimal = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if decimal == "," else ",", "").replace(",", ".")
    elif "," in s:
        tail = s.rsplit(",", 1)[1]
        thousands = s.count(",") > 1 or (len(tail) == 3 and decimal_comma is not True)
        s = s.replace(",", "") if thousands else s.replace(",", ".")
    elif s.count(".") > 1 or (len(s.rsplit(".", 1)[-1]) == 3 and "." in s and decimal_comma):
        s = s.replace(".", "")
    try:
        result = Decimal(s)
    except InvalidOperation:
        raise ValueError(text) from None
    if not result.is_finite():
        raise ValueError(text)
    return -result if negative else result


_ISO = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})")
_NUMERIC = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})")
_COMPACT = re.compile(r"^(\d{4})(\d{2})(\d{2})$")
_TIME = re.compile(
    r"[\sT,]+\d{1,2}[:.]\d{2}([:.]\d{2}(\.\d+)?)?(\s*[ap]\.?m\.?)?(\s*(z|[+-]\d{2}:?\d{2}|[a-z]{3,4}))?$",
    re.I,
)
_TEXT_FORMATS = ("%d %b %Y", "%d %B %Y", "%b %d %Y", "%B %d %Y", "%a %d %b %Y", "%A %d %B %Y")


def day_first(values: list[Any]) -> bool:
    """Whether numeric dates in this file are day/month (the Tricount default) or month/day."""
    for v in values:
        m = _NUMERIC.match(str(v or "").strip())
        if m and int(m.group(1)) > 12:
            return True
        if m and int(m.group(2)) > 12:
            return False
    return True


def parse_date(value: Any, dmy: bool = True) -> date | None:
    """A cell -> date, or None if blank. Raises ValueError for anything unreadable."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        return date(1899, 12, 30) + timedelta(days=int(value))  # an Excel serial day
    text = str(value).strip()
    if not text:
        return None
    try:
        if m := _ISO.match(text):
            y, mo, d = (int(g) for g in m.groups())
            return date(y, mo, d)
        if m := _COMPACT.match(text):
            y, mo, d = (int(g) for g in m.groups())
            return date(y, mo, d)
        if m := _NUMERIC.match(text):
            a, b, y = (int(g) for g in m.groups())
            y = y + 2000 if y < 100 else y
            return date(y, b, a) if dmy else date(y, a, b)
    except ValueError:
        raise ValueError(text) from None
    bare = re.sub(r"[,.]", " ", _TIME.sub("", text))
    bare = re.sub(r"\s+", " ", bare).strip()
    for fmt in _TEXT_FORMATS:
        try:
            return datetime.strptime(bare, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def _currency(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    return _SYMBOLS.get(text.lower(), text.upper())


def _kind(value: Any) -> tuple[str, str | None]:
    """(kind, note). Unknown types import as expenses, with a note."""
    text = str(value or "").strip()
    n = _norm(text)
    if not n:
        return "expense", None
    for kind, names in _KINDS.items():
        if n in names:
            return kind, None
    return "expense", f"Type “{text}” isn't one we know, so it's imported as an expense"


def _same_person(people: dict[str, str], name: str) -> str:
    """The first spelling seen of a name, so "rahul s" in the payer column joins "Rahul S"."""
    return people.setdefault(_norm(name), name.strip())


# ---- the whole file ----


def parse(content: bytes, file_name: str) -> ParsedFile:
    table = read_table(content, file_name)
    at, cols = _find_header(table)
    title = next(
        (str(c).strip() for r in table[:at] for c in r if c is not None and str(c).strip()), None
    )
    body = table[at + 1 :]

    def cell(row: list[Any], key: str) -> Any:
        i = cols.fixed.get(key)
        return row[i] if i is not None and i < len(row) else None

    def at_col(row: list[Any], i: int) -> Any:
        return row[i] if i < len(row) else None

    amount_cells = [
        at_col(r, i)
        for r in body
        for i in [
            *(cols.fixed[k] for k in ("amount", "amount_default") if k in cols.fixed),
            *(i for i, _ in cols.shares + cols.paid_by),
        ]
    ]
    comma = decimal_style(amount_cells)
    dmy = day_first([cell(r, "date") for r in body])

    people: dict[str, str] = {}
    for _, name in cols.shares + cols.paid_by:
        _same_person(people, name)

    rows: list[RawRow] = []
    for n, r in enumerate(body, start=at + 2):
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        problems: list[str] = []
        notes: list[str] = []

        def amount(value: Any, what: str) -> Decimal | None:
            try:
                return parse_amount(value, comma)
            except ValueError:
                problems.append(f"{what} “{value}” isn't a number we can read")  # noqa: B023
                return None

        date_cell = cell(r, "date")
        try:
            day = parse_date(date_cell, dmy)
        except ValueError:
            day = None
            problems.append(f"The date “{date_cell}” isn't one we can read")
        kind, note = _kind(cell(r, "type"))
        if note:
            notes.append(note)
        payer_cell = str(cell(r, "payer") or "").strip()
        if not payer_cell and not str(date_cell or "").strip() and "payer" in cols.fixed:
            continue  # a totals or balance line under the table, not a transaction
        payer = _same_person(people, payer_cell) if payer_cell else None
        shares = {
            _same_person(people, name): v
            for i, name in cols.shares
            if (v := amount(at_col(r, i), f"{name}'s share")) is not None
        }
        paid_by = {
            _same_person(people, name): v
            for i, name in cols.paid_by
            if (v := amount(at_col(r, i), f"What {name} paid")) is not None
        }
        category = (
            str(cell(r, "custom_category") or "").strip() or str(cell(r, "category") or "").strip()
        )
        rows.append(
            RawRow(
                line=n,
                title=str(cell(r, "title") or "").strip(),
                date=day,
                date_text=str(date_cell or ""),
                kind=kind,
                kind_text=str(cell(r, "type") or "").strip(),
                currency=_currency(cell(r, "currency")),
                amount=amount(cell(r, "amount"), "The amount"),
                amount_default=amount(cell(r, "amount_default"), "The converted amount"),
                payer=payer,
                paid_by=paid_by,
                shares=shares,
                category=None if _norm(category) in ("", "uncategorized") else category,
                problems=problems,
                notes=notes,
            )
        )
    if not rows:
        raise ParseError("This file has Tricount's columns but no expenses in it.")
    return ParsedFile(rows=rows, people=list(people.values()), title=title)
